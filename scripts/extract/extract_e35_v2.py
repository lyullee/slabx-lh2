"""
Extract the PRESLHY E3.5 far-field concentrations -- second version.

Two defects in v1, both found by comparing the output against D3.6:

1. **The weather was averaged over the whole day.**  `LocalWeather` holds a
   5-minute record running for hours, so trials 10-15 all came out with the
   same 2.21 m/s and trials 19-25 all with 1.47 m/s.  This version windows it
   to each trial's own release, using the clock times in the `Draeger` sheet.

2. **The flow meter is not usable at 1 bar.**  D3.6 §"Mass flow rates" says
   so: at low pressure and through the open pipe the meter saw two-phase flow
   with a high void fraction and did not give reliable output.  Checked
   against D3.6 Table 4, the measured peak matches at 5 bar (trial 12: 99.9
   against 90-100 g/s; trial 10: 285 against 298) and is a third of it at
   1 bar (trial 17: 30.0 against 104-107).  This version reports the measured
   value and the Table 4 value side by side and marks which to trust, rather
   than silently picking one.

Also recorded, because it changes the diagnostics: D3.6 §2.1.6 puts the
near-field weather station at **1.5 m**, not the 2 m assumed earlier, and the
far-field station on a 3 m stand about 20 m downwind.

    pip install openpyxl
    python extract_e35_v2.py "C:\\...\\10.35097-1481"
    python extract_e35_v2.py "C:\\...\\10.35097-1481" --timeseries 13,20,16,12,17
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
from pathlib import Path

# --- D3.6 Table A4: instrument serial -> (stand, height above ground) -------
SENSORS = {
    "ARBC0893": (1, 0.5), "ARLF0195": (1, 1.5), "ARDD0512": (1, 2.5),
    "ARFF0195": (2, 0.5), "ARFK0515": (2, 1.5), "ARFK0724": (2, 2.5),
    "ARJB0150": (3, 0.5), "ARFJ0274": (3, 1.5), "ARFK0235": (3, 2.5),
    "ARCC0390": (4, 0.5), "AREH0319": (4, 1.5), "AREL0225": (4, 2.5),
    "AREJ0521": (5, 0.5), "ARFK0648": (5, 1.5), "ARAK0249": (5, 2.5),
    "ARHK0084": (6, 0.5), "ARHK0130": (6, 1.5), "ARFK0590": (6, 2.5),
    "ARAM0219": (7, 0.5), "ARFJ0366": (7, 1.5), "ARHM0237": (7, 2.5),
    "ARLF0431": (8, 0.5), "ARBC0822": (8, 1.5), "ARBD0399": (8, 2.5),
    "AREL0307": (9, 0.5), "ARZD0946": (9, 1.5), "ARLF0193": (9, 2.5),
    "ARLF0465": (10, 0.5), "ARCC0403": (10, 1.5), "ARFK0507": (10, 2.5),
}

TRIAL_TEST = {
    1: "3.5.3", 2: "3.5.1", 3: "3.5.1", 4: "3.5.2", 5: "3.5.3", 6: "3.5.7",
    7: "3.5.8", 8: "3.5.8", 9: "3.5.9", 10: "3.5.10", 11: "3.5.11",
    12: "3.5.12", 13: "3.5.17", 14: "3.5.16", 15: "3.5.18", 16: "3.5.4",
    17: "3.5.5", 18: "3.5.6", 19: "3.5.4", 20: "3.5.5", 21: "3.5.6",
    22: "3.5.13", 23: "3.5.14", 24: "3.5.15", 25: "3.5.13",
}
TEST_COND = {
    "3.5.1": ("horizontal", 0.5, 25.4, 1), "3.5.2": ("horizontal", 0.5, 12, 1),
    "3.5.3": ("horizontal", 0.5, 6, 1),    "3.5.4": ("horizontal", 1.5, 25.4, 1),
    "3.5.5": ("horizontal", 1.5, 12, 1),   "3.5.6": ("horizontal", 1.5, 6, 1),
    "3.5.7": ("vertical_up", 0.5, 12, 1),  "3.5.8": ("vertical_down", 0.5, 12, 1),
    "3.5.9": ("horizontal_obstruction", 0.5, 12, 1),
    "3.5.10": ("horizontal", 0.5, 25.4, 5), "3.5.11": ("horizontal", 0.5, 12, 5),
    "3.5.12": ("horizontal", 0.5, 6, 5),   "3.5.13": ("horizontal", 1.5, 25.4, 5),
    "3.5.14": ("horizontal", 1.5, 12, 5),  "3.5.15": ("horizontal", 1.5, 6, 5),
    "3.5.16": ("vertical_up", 0.5, 12, 5), "3.5.17": ("vertical_down", 0.5, 12, 5),
    "3.5.18": ("horizontal_obstruction", 0.5, 12, 5),
}
#: D3.6 Table 4, g/s.  None where the report itself says the value is unknown.
TABLE4_FLOW = {(6, 5): 95.0, (12, 5): 265.0, (25.4, 5): 298.0,
               (6, 1): None, (12, 1): 105.5, (25.4, 1): 139.5}

TRIAL_RE = re.compile(r"trial[_\s]*(\d+)", re.I)
H2_COL = re.compile(r"^([A-Z]{4}\d{4})_.*_dlH2percent$", re.I)
O2_COL = re.compile(r"^([A-Z]{4}\d{4})_.*_dlO2percent$", re.I)
SAT = 3.90        # Draeger %vol over-range; D3.6 warns about 4 % saturation


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def clock(v):
    """Seconds since midnight from whatever the sheet stores as a time."""
    if v is None:
        return None
    if isinstance(v, dt.time):
        return v.hour*3600 + v.minute*60 + v.second + v.microsecond/1e6
    if isinstance(v, dt.datetime):
        return v.hour*3600 + v.minute*60 + v.second + v.microsecond/1e6
    m = re.match(r"\s*(\d{1,2}):(\d{2}):(\d{2}(?:\.\d+)?)", str(v))
    if not m:
        return None
    return int(m.group(1))*3600 + int(m.group(2))*60 + float(m.group(3))


def read_sheet(wb, name):
    if name not in wb.sheetnames:
        return None, None
    rows = wb[name].iter_rows(values_only=True)
    try:
        headers = [str(c) if c is not None else "" for c in next(rows)]
    except StopIteration:
        return None, None
    return headers, list(rows)


def col_index(headers, needle):
    for i, h in enumerate(headers):
        if needle.lower() in h.lower():
            return i
    return None


def window_stats(headers, rows, needle, lo, hi, time_needle="time"):
    """max/mean of a column, restricted to rows whose clock is in [lo, hi]."""
    i = col_index(headers, needle)
    t = col_index(headers, time_needle)
    if i is None:
        return "", "", 0
    vals = []
    for r in rows:
        if t is not None and lo is not None:
            c = clock(r[t]) if t < len(r) else None
            if c is None or not (lo <= c <= hi):
                continue
        v = num(r[i]) if i < len(r) else None
        if v is not None:
            vals.append(v)
    if not vals:
        return "", "", 0
    return max(vals), sum(vals)/len(vals), len(vals)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("root")
    ap.add_argument("--timeseries", default="")
    ap.add_argument("--pad", type=float, default=300.0,
                    help="seconds of weather record to keep either side")
    args = ap.parse_args()

    try:
        from openpyxl import load_workbook
    except ImportError:
        print("pip install openpyxl")
        return 2

    root = Path(args.root).expanduser().resolve()
    files = sorted(root.rglob("trial_*alldata.xlsx"))
    if not files:
        print(f"no trial_*alldata.xlsx under {root}")
        return 2
    want_ts = {int(t) for t in args.timeseries.split(",") if t.strip()}
    here = Path(__file__).resolve().parent
    far, cond, ts = [], [], []

    for path in files:
        m = TRIAL_RE.search(path.name)
        if not m:
            continue
        trial = int(m.group(1))
        wb = load_workbook(path, read_only=True, data_only=True)
        dh, dr = read_sheet(wb, "Draeger")

        # the trial's own clock window, from the Draeger record
        lo = hi = None
        if dh:
            ti = col_index(dh, "dltime")
            if ti is not None:
                cs = [c for c in (clock(r[ti]) for r in dr if ti < len(r))
                      if c is not None]
                if cs:
                    lo, hi = min(cs), max(cs)

        h2i, o2i = {}, {}
        for i, h in enumerate(dh or []):
            g = H2_COL.match(h)
            if g:
                h2i[g.group(1).upper()] = i
                continue
            g = O2_COL.match(h)
            if g:
                o2i[g.group(1).upper()] = i

        for serial, i in h2i.items():
            vals = [v for v in (num(r[i]) for r in dr if i < len(r))
                    if v is not None]
            if not vals:
                continue
            peak = max(vals)
            active = [v for v in vals if v > 0.1*peak] if peak > 0 else []
            j = o2i.get(serial)
            o2 = ([v for v in (num(r[j]) for r in dr if j < len(r))
                   if v is not None] if j is not None else [])
            stand, height = SENSORS.get(serial, ("", ""))
            far.append({
                "trial": trial, "serial": serial, "stand": stand,
                "height_m": height, "n": len(vals),
                "h2_max": peak,
                "saturated": int(peak >= SAT),
                "n_at_ceiling": sum(1 for v in vals if v >= SAT),
                "h2_mean_all": sum(vals)/len(vals),
                "h2_mean_active": (sum(active)/len(active)) if active else 0.0,
                "n_active": len(active),
                "o2_min": min(o2) if o2 else "",
            })

        wh, wr = read_sheet(wb, "LocalWeather")
        lo_w = None if lo is None else lo - args.pad
        hi_w = None if hi is None else hi + args.pad
        u_max, u_mean, n_w = window_stats(wh, wr, "WindSpeed", lo_w, hi_w) if wh else ("","",0)
        g_max, _, _ = window_stats(wh, wr, "Gust", lo_w, hi_w) if wh else ("","",0)
        t_max, t_mean, _ = window_stats(wh, wr, "OutdoorTemperature", lo_w, hi_w) if wh else ("","",0)
        _, rh_mean, _ = window_stats(wh, wr, "OutdoorHumidity", lo_w, hi_w) if wh else ("","",0)

        fh, fr = read_sheet(wb, "Flowmeter")
        q_max, q_mean, n_q = window_stats(fh, fr, "mass", None, None) if fh else ("","",0)

        test = TRIAL_TEST.get(trial, "")
        orient, rel_h, orifice, bar = TEST_COND.get(test, ("", "", "", ""))
        t4 = TABLE4_FLOW.get((orifice, bar))
        # the meter is trustworthy at 5 bar and not at 1 bar (D3.6)
        use = "meter_peak" if bar == 5 else ("table4" if t4 else "unknown")
        cond.append({
            "trial": trial, "test_no": test, "orientation": orient,
            "release_height_m": rel_h, "orifice_mm": orifice, "tanker_barg": bar,
            "wind_mean_ms": u_mean, "wind_max_ms": u_max, "gust_max_ms": g_max,
            "wind_ref_height_m": 1.5,          # D3.6 section 2.1.6
            "n_weather_samples": n_w,
            "T_mean_C": t_mean, "RH_mean_pct": rh_mean,
            "flow_meter_peak_gs": q_max, "flow_meter_mean_gs": q_mean,
            "flow_table4_gs": t4 if t4 is not None else "",
            "flow_source": use,
            "flow_used_gs": (q_max if use == "meter_peak"
                             else (t4 if use == "table4" else "")),
            "window_start_s": lo if lo is not None else "",
            "window_end_s": hi if hi is not None else "",
            "n_draeger_sensors": len(h2i), "file": path.name,
        })

        if trial in want_ts and dh:
            ti = col_index(dh, "dltime")
            cols = [(i, s) for s, i in h2i.items()]
            for r in dr:
                tv = clock(r[ti]) if ti is not None and ti < len(r) else ""
                for i, serial in cols:
                    v = num(r[i]) if i < len(r) else None
                    if v is None:
                        continue
                    stand, height = SENSORS.get(serial, ("", ""))
                    ts.append({"trial": trial, "t_s": tv, "serial": serial,
                               "stand": stand, "height_m": height,
                               "h2_percent": v})
        wb.close()
        print(f"  trial {trial:>2}  window {lo}-{hi}  weather rows {n_w}")

    def dump(name, rows):
        if not rows:
            return
        p = here / name
        with p.open("w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
            w.writeheader(); w.writerows(rows)
        print(f"wrote {p}  ({len(rows)} rows)")

    dump("lh2_e35_farfield_v2.csv", far)
    dump("lh2_e35_conditions_v2.csv", cond)
    dump("lh2_e35_timeseries.csv", ts)

    unknown = sorted({r["serial"] for r in far if r["stand"] == ""})
    if unknown:
        print("\n!! serials not in Table A4 -- report, do not guess:", unknown)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
